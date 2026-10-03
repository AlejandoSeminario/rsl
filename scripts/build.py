"""Genera la base de datos Excel, el formulario de extracción y los conteos PRISMA."""
import json, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

recs = json.load(open('../data/recs2.json'))
s2q = json.load(open('../data/s2_all.json'))['query']

# Decisión de elegibilidad (lectura del registro completo) sobre los 95 candidatos.
INCL = ("R0212 R0479 R0556 R0185 R0201 R0226 R0236 R0264 R0473 R0477 R0488 R0779 R0051 R0070 R0286 "
        "R0293 R0423 R0445 R0492 R0623 R0817 R0852 R0072 R0076 R0210 R0292 R0400 R0405 R0416 R0460 "
        "R0516 R0528 R0617 R0697 R0333 R0671 R0431 R0538 R0575 R0591 R0615 R0754 R0862 R0020 R0588 R0688").split()
EXCL_FT = {
 'CE4: estudio secundario (revisión), se usa solo como antecedente': "R0339 R0439 R0535 R0553 R0594 R0650 R0685",
 'CE2: robot sin función social/asistencial (conducción, manipulación industrial, doméstica o logística)': "R0016 R0147 R0541 R0703 R0790 R0424 R0219 R0313 R0539 R0699 R0001 R0347",
 'CE3: sin robot físico (agente virtual, interfaz, dataset o reto de evaluación)': "R0376 R0047 R0340 R0410 R0629 R0645 R0283 R0158 R0402 R0356 R0018 R0295 R0331 R0302 R0363 R0216 R0253 R0269 R0686 R0744 R0198 R0620 R0526 R0842 R0562 R0288 R0095 R0595 R0200 R0747",
}
excl_ft = {i: k for k, v in EXCL_FT.items() for i in v.split()}

cand = set()
for r in recs:
    r['etapa'] = ''; r['decision'] = ''; r['motivo'] = ''
    if r['dup_of']:
        r['etapa'] = '1. Eliminación de duplicados'; r['decision'] = 'Excluido'; r['motivo'] = f"Duplicado de {r['dup_of']}"; continue
    if r.get('filter'):
        r['etapa'] = '2. Filtros de tipo/resumen'; r['decision'] = 'Excluido'; r['motivo'] = r['filter']; continue
    ok = r['c_soc'] and r['c_gen'] and not r['nonsoc'] and r['multi'] and r['empir'] and r['n_out'] >= 3
    if not ok:
        r['etapa'] = '3. Cribado título/resumen'; r['decision'] = 'Excluido'
        r['motivo'] = ('CE2: no aborda robótica social/HRI' if (not r['c_soc'] or r['nonsoc']) else
                       'CI2 no cumplido: no integra IA generativa' if not r['c_gen'] else
                       'CI3 no cumplido: no hay interacción multimodal' if not r['multi'] else
                       'CI4 no cumplido: sin evaluación empírica o resultados relevantes')
        continue
    cand.add(r['id']); r['etapa'] = '4. Elegibilidad (registro completo)'
    if r['id'] in INCL: r['decision'] = 'Incluido'; r['motivo'] = 'Cumple CI1–CI5'
    else: r['decision'] = 'Excluido'; r['motivo'] = excl_ft.get(r['id'], 'CE3: sin robot físico (agente virtual, interfaz, dataset o reto de evaluación)')
assert set(INCL) <= cand, set(INCL) - cand

uniq = [r for r in recs if not r['dup_of']]
c = dict(
  s2=sum(r['db']=='Semantic Scholar' for r in recs), cr=sum(r['db']=='Crossref' for r in recs),
  total=len(recs), dup=sum(bool(r['dup_of']) for r in recs),
  f_type=sum(r['motivo'].startswith('CE1') and r['etapa'].startswith('2') for r in recs),
  f_abs=sum(r['motivo'].startswith('CE5') for r in recs),
  screened=sum(r['etapa'] in ('3. Cribado título/resumen','4. Elegibilidad (registro completo)') for r in recs),
  ex_screen=sum(r['etapa']=='3. Cribado título/resumen' for r in recs),
  elig=len(cand), incl=len(INCL))
from collections import Counter
c['ex_screen_by'] = Counter(r['motivo'] for r in recs if r['etapa']=='3. Cribado título/resumen')
c['ex_ft_by'] = Counter(r['motivo'] for r in recs if r['etapa'].startswith('4') and r['decision']=='Excluido')
c['incl_years'] = Counter(r['year'] for r in recs if r['id'] in INCL)
c['incl_types'] = Counter(r['type'] for r in recs if r['id'] in INCL)
json.dump({k: (dict(v) if isinstance(v, Counter) else v) for k, v in c.items()}, open('../data/prisma_counts.json','w'), indent=1, ensure_ascii=False, default=str)
print(json.dumps({k: (dict(v) if isinstance(v, Counter) else v) for k, v in c.items()}, indent=1, ensure_ascii=False, default=str))

H = Font(bold=True, color='FFFFFF'); FILL = PatternFill('solid', fgColor='1F4E78')
def header(ws, cols, widths):
    ws.append(cols)
    for i, cell in enumerate(ws[1]):
        cell.font = H; cell.fill = FILL; cell.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[cell.column_letter].width = widths[i] if i < len(widths) else 18
    ws.freeze_panes = 'A2'

wb = Workbook()
ws = wb.active; ws.title = 'Ecuaciones'
header(ws, ['Base de datos', 'Campo de búsqueda', 'Ecuación aplicada', 'Filtros', 'Fecha de ejecución', 'Registros recuperados'], [18, 22, 110, 40, 14, 12])
ws.append(['Semantic Scholar (Academic Graph API, /paper/search/bulk)', 'Título y resumen', s2q, 'Años 2020–2026', '2026-10-03', c['s2']])
ws.append(['Crossref (REST API /works)', 'Título y resumen (equación aplicada sobre los metadatos)',
           'Consultas P×I: {"social robot" | "social robotics" | "socially assistive robot" | "companion robot" | "human-robot interaction"} × {"large language model" | "generative AI" | "vision-language model" | "foundation model" | "GPT" | "ChatGPT"}; luego se conservaron solo los registros que cumplen la misma ecuación booleana en título+resumen',
           'from-pub-date 2020-01-01, until-pub-date 2026-12-31', '2026-10-03', c['cr']])
ws.append(['Scopus (equivalente, para replicación)', 'TITLE-ABS-KEY',
           'TITLE-ABS-KEY ( "social robot*" OR "social robotics" OR "socially assistive robot*" OR "companion robot*" OR "human-robot interaction" ) AND TITLE-ABS-KEY ( "large language model*" OR "generative AI" OR "generative artificial intelligence" OR "vision-language model*" OR "foundation model*" OR gpt OR chatgpt OR "multimodal LLM" ) AND PUBYEAR > 2019 AND PUBYEAR < 2027 AND ( LIMIT-TO ( DOCTYPE , "ar" ) OR LIMIT-TO ( DOCTYPE , "cp" ) ) AND ( LIMIT-TO ( LANGUAGE , "English" ) OR LIMIT-TO ( LANGUAGE , "Spanish" ) )',
           'Igual que arriba', 'Pendiente', ''])
ws.append(['Web of Science (equivalente, para replicación)', 'TS (Topic)',
           'TS=("social robot*" OR "social robotics" OR "socially assistive robot*" OR "companion robot*" OR "human-robot interaction") AND TS=("large language model*" OR "generative AI" OR "generative artificial intelligence" OR "vision-language model*" OR "foundation model*" OR GPT OR ChatGPT OR "multimodal LLM")',
           'PY=2020-2026; DT=(Article OR Proceedings Paper); LA=(English OR Spanish)', 'Pendiente', ''])
for row in ws.iter_rows(min_row=2):
    for cell in row: cell.alignment = Alignment(wrap_text=True, vertical='top')

ws = wb.create_sheet('Registros')
header(ws, ['ID', 'Base de datos', 'Título', 'Autores', 'Año', 'Fuente', 'Tipo', 'DOI', 'Resumen', 'Etapa PRISMA', 'Decisión', 'Motivo / criterio'],
       [8, 15, 60, 35, 6, 35, 15, 28, 60, 26, 10, 50])
for r in recs:
    ws.append([r['id'], r['db'], r['title'], r['authors'][:500], r['year'], r['source'], r['type'], r['doi'], r['abstract'][:3000], r['etapa'], r['decision'], r['motivo']])
ws.auto_filter.ref = ws.dimensions

ws = wb.create_sheet('Incluidos')
header(ws, ['N.º', 'ID', 'Autores', 'Año', 'Título', 'Fuente', 'Tipo', 'DOI'], [5, 8, 40, 6, 70, 40, 15, 30])
inc = sorted([r for r in recs if r['id'] in INCL], key=lambda r: (r['authors'].split(';')[0].split()[-1] if r['authors'] else r['title']))
for n, r in enumerate(inc, 1):
    ws.append([n, r['id'], r['authors'][:300], r['year'], r['title'], r['source'], r['type'], r['doi']])

ws = wb.create_sheet('PRISMA')
header(ws, ['Etapa', 'n'], [70, 10])
for k, v in [('Registros identificados en Semantic Scholar', c['s2']), ('Registros identificados en Crossref', c['cr']),
             ('Total identificados', c['total']), ('Duplicados eliminados', c['dup']),
             ('Excluidos por tipo de documento / preprint (CE1)', c['f_type']), ('Excluidos sin resumen (CE5)', c['f_abs']),
             ('Registros cribados por título y resumen', c['screened']), ('Excluidos en cribado', c['ex_screen']),
             ('Informes evaluados para elegibilidad', c['elig']), ('Excluidos en elegibilidad', c['elig'] - c['incl']),
             ('Estudios incluidos en la RSL', c['incl'])]:
    ws.append([k, v])
wb.save('../Base_de_datos.Seminario-Ordaya.xlsx')

# ---------- Formulario de extracción ----------
wb = Workbook(); ws = wb.active; ws.title = 'Formulario'
cols = ['N.º', 'ID', 'Referencia (APA)', 'DOI', 'Año', 'País / institución', 'Tipo de estudio (diseño)',
        'Plataforma robótica', 'Población / contexto (salud, educación, compañía…)', 'N.º de participantes',
        'Modelo generativo usado (LLM/VLM)', 'Despliegue (nube / local / borde)', 'Modalidades de entrada (voz, texto, imagen, video, gesto)',
        'Modalidades de salida (habla, gesto, mirada, expresión facial, movimiento)', 'Arquitectura / método de fusión sensorial',
        'Traducción lenguaje-acción (sí/no; cómo)', 'Latencia reportada (s)', 'Estrategia de mitigación de latencia',
        'Sincronización verbal-no verbal (cómo)', 'Errores ASR / robustez reportada', 'Limitaciones cognitivas (ToM, razonamiento espacial, alucinaciones)',
        'Riesgos éticos identificados (antropomorfización, manipulación, sesgos)', 'Salvaguardas / gobernanza ética propuesta',
        'Métricas de evaluación', 'Resultados principales', 'Limitaciones declaradas por los autores', 'Dimensión(es) que aborda (técnica/cognitiva/ética)', 'Observaciones del revisor']
header(ws, cols, [5, 8, 60, 28, 6] + [24]*30)
dv_dep = DataValidation(type='list', formula1='"Nube,Local,Borde (edge),Híbrido,No reportado"', allow_blank=True)
dv_dim = DataValidation(type='list', formula1='"Técnica,Cognitiva,Ética,Técnica+Cognitiva,Técnica+Ética,Cognitiva+Ética,Las tres"', allow_blank=True)
dv_yn = DataValidation(type='list', formula1='"Sí,No,Parcial,No reportado"', allow_blank=True)
for dv in (dv_dep, dv_dim, dv_yn): ws.add_data_validation(dv)
def apa(r):
    a = [x.strip() for x in r['authors'].split(';') if x.strip()]
    def fmt(n):
        p = n.split(); return (p[-1] + ', ' + ' '.join(q[0] + '.' for q in p[:-1])) if len(p) > 1 else n
    a = [fmt(x) for x in a]
    au = (', '.join(a[:-1]) + ', & ' + a[-1]) if len(a) > 1 else (a[0] if a else '')
    if len(a) > 20: au = ', '.join(a[:19]) + ', … ' + a[-1]
    return f"{au} ({r['year']}). {r['title']}. {r['source']}." + (f" https://doi.org/{r['doi']}" if r['doi'] else '')
for n, r in enumerate(inc, 1):
    t = (r['title'] + ' ' + r['abstract']).lower()
    mods_in = ', '.join(m for m, p in [('voz', r'speech|voice|spoken|asr'), ('texto', r'text'), ('imagen/visión', r'vision|visual|image|camera'), ('video', r'video'), ('gesto/postura', r'gesture|posture|pose'), ('mirada', r'gaze'), ('señales fisiológicas', r'physiolog|biosens|eeg|heart')] if re.search(p, t))
    ws.append([n, r['id'], apa(r), r['doi'], r['year'], '', '', '', '', '', '', '', mods_in])
    row = ws.max_row
    dv_dep.add(f'L{row}'); dv_yn.add(f'P{row}'); dv_dim.add(f'AA{row}')
for row in ws.iter_rows(min_row=2):
    for cell in row: cell.alignment = Alignment(wrap_text=True, vertical='top')
ws2 = wb.create_sheet('Instrucciones')
for line in ['Formulario de extracción de datos de la RSL "Integración de IA Generativa Multimodal en la HRI para Robótica Social (2020–2026)".',
             'Una fila por estudio incluido. Las columnas siguen las dimensiones técnica, cognitiva y ética definidas en el objetivo y en la pregunta PICO.',
             'La columna "Modalidades de entrada" está precargada a partir del título y el resumen; verificarla al leer el texto completo.',
             'Usar "No reportado" cuando el estudio no informe el dato. Las latencias se registran en segundos (media ± DE si se reporta).',
             'Dos revisores extraen los datos de forma independiente; las discrepancias se resuelven por consenso.']:
    ws2.append([line])
ws2.column_dimensions['A'].width = 140
wb.save('../Formulario.Seminario-Ordaya.xlsx')
json.dump([dict(id=r['id'], apa=apa(r), authors=r['authors'], year=r['year'], title=r['title'], source=r['source'], doi=r['doi'], type=r['type']) for r in inc], open('../data/incluidos.json', 'w'), ensure_ascii=False, indent=1)
